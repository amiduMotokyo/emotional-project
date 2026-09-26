"""Build paper assets from archived results and export the single Markdown source."""
from pathlib import Path
import csv,json,hashlib,re,sys,shutil
from copy import deepcopy
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from C.scripts.export_round2_paper import inline,element,text
SOURCE=ROOT/'docs/论文稿/问题二.md'
ASSETS=ROOT/'docs/C/paper/assets/q2_integrated'
RUN=ROOT/'C/outputs/q2_optimization_local_myenv_v1'
TEAM=ROOT/'C/outputs/q2_accuracy_screen_20260925'
plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei'],'axes.unicode_minus':False,'font.size':10})

def save(fig,name):
    fig.savefig(ASSETS/f'{name}.png',dpi=240,bbox_inches='tight')
    fig.savefig(ASSETS/f'{name}.svg',bbox_inches='tight');plt.close(fig)

def assets():
    ASSETS.mkdir(parents=True,exist_ok=True)
    grid=json.loads((RUN/'validation_grid.json').read_text(encoding='utf-8'))['rows']
    rows=[r for r in grid if r['group']=='E0' and r['seed']==20260924]
    assert len(rows)==64
    missing=[r for r in rows if r['case']!='clean']
    avg={k:float(np.mean([r['coherent'][k] for r in missing])) for k in ['accuracy','macro_f1','mae','pearson']}
    modes=['text','audio','vision','text+audio','text+vision','audio+vision','text+audio+vision']
    rates=[10,30,50];positions=['start','middle','end']
    def metric(mode=None,rate=None,pos=None):
        rr=[]
        for r in missing:
            m,rate_s,p=r['case'].replace('.npz','').split('_');m=m.replace('-','+')
            if (mode is None or m==mode) and (rate is None or int(rate_s)==rate) and (pos is None or p==pos):rr.append(r['coherent']['macro_f1'])
        assert rr,(mode,rate,pos)
        return np.mean(rr)
    heat=np.array([[metric(m,r) for r in rates] for m in modes])
    fig,axs=plt.subplots(1,2,figsize=(10,4.7),gridspec_kw={'width_ratios':[1.1,1]})
    im=axs[0].imshow(heat,cmap='Blues',vmin=.50,vmax=.62,aspect='auto')
    axs[0].set_yticks(range(7),['文本','音频','视觉','文本＋音频','文本＋视觉','音频＋视觉','三模态']);axs[0].set_xticks(range(3),['10%','30%','50%'])
    for i in range(7):
        for j in range(3):axs[0].text(j,i,f'{heat[i,j]:.3f}',ha='center',va='center',color='white' if heat[i,j]>.575 else 'black')
    axs[0].set_title('模态组合 × 缺失比例');fig.colorbar(im,ax=axs[0],shrink=.7,label='Macro-F1')
    for p,l,c in zip(positions,['起始','中部','末尾'],['#355C7D','#6C9A8B','#D8933D']):axs[1].plot(rates,[metric(rate=r,pos=p) for r in rates],'-o',label=l,color=c)
    axs[1].set(xlabel='名义缺失比例（%）',ylabel='Macro-F1',title='位置与长度的联合影响',xticks=rates);axs[1].legend();axs[1].grid(alpha=.2);fig.tight_layout();save(fig,'missingness')
    cm=np.loadtxt(RUN/'paper/validation/E0_20260924_confusion.csv',delimiter=',',skiprows=1,usecols=(1,2,3),dtype=int)
    fig,ax=plt.subplots(figsize=(6.2,4.5));ax.imshow(cm/cm.sum(1)[:,None],cmap='Blues',vmin=0,vmax=1)
    for i in range(3):
        for j in range(3):ax.text(j,i,f'{cm[i,j]}\n({cm[i,j]/cm[i].sum():.1%})',ha='center',va='center',color='white' if cm[i,j]/cm[i].sum()>.55 else 'black')
    ax.set(xticks=range(3),yticks=range(3),xticklabels=['负向','中性','正向'],yticklabels=['负向','中性','正向'],xlabel='预测类别',ylabel='真实类别');fig.tight_layout();save(fig,'confusion')
    fig,ax=plt.subplots(figsize=(10,3.8));xx=np.arange(5)
    ax.bar(xx-.17,[61.3095,60.4396,60.8516,61.0806,62.7747],.34,label='原Fusion A0',color='#355C7D');ax.bar(xx+.17,[61.3095,61.5842,62.1795,62.4084,62.5],.34,label='动态时序 B2',color='#6C9A8B')
    ax.set(xticks=xx,xticklabels=['FP32均值','int8均值','折外校准均值','全valid拟合均值','预固定单模型'],ylim=(58,65),ylabel='官方valid ACC（%）');ax.legend(ncol=2);ax.grid(axis='y',alpha=.2);fig.tight_layout();save(fig,'precision')
    fig,ax=plt.subplots(figsize=(11,3.5));ax.set(xlim=(0,11),ylim=(0,3.5));ax.axis('off')
    boxes=[(.1,2.3,2.3,.8,'文本词元＋连续缺失\n先替换UNK，再编码'),(.1,1.2,2.3,.8,'音频74维／视觉35维\n训练统计标准化'),(3,2.3,2,.8,'MiniLM序列\n50×384'),(3,1.2,2,.8,'三路投影＋\n掩码注意力池化'),(5.6,1.7,2.1,1.1,'可用比例＋门控\n三路聚合与加权融合'),(8.4,1.7,2.4,1.1,'联合分类／回归\n符号一致性输出')]
    for x,y,w,h,s in boxes:ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.06',fc='#EDF2F6',ec='#355C7D'));ax.text(x+w/2,y+h/2,s,ha='center',va='center')
    for a,b in [((2.5,2.7),(2.9,2.7)),((2.5,1.6),(2.9,1.6)),((4,2.25),(4,2.05)),((5.05,1.7),(5.5,2.1)),((7.8,2.25),(8.3,2.25))]:ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'->','color':'#355C7D'})
    ax.text(5.5,.45,'训练：加权交叉熵＋0.8 SmoothL1；验证：原始条件＋63种连续缺失条件',ha='center');save(fig,'architecture')
    pred=list(csv.DictReader((RUN/'attachment3_predictions.csv').open(encoding='utf-8-sig')))
    assert len(pred)==len({r['sample_id'] for r in pred})==30
    table='| 样本编号 | 极性 | 强度 |\n|---|---|---:|\n'+'\n'.join(f"| {r['sample_id']} | {dict(Negative='负向',Neutral='中性',Positive='正向')[r['polarity']]} | {float(r['intensity']):.6f} |" for r in pred)
    discussion='对七种模态组合和三个位置平均，10%、30%、50%缺失时Macro-F1分别为'+ '、'.join(f'{metric(rate=r):.4f}' for r in rates)+'。对全部比例与模态组合平均，起始、中部、末尾缺失时分别为'+'、'.join(f'{metric(pos=p):.4f}' for p in positions)+'；上述位置比较来自相同模型和同一验证集，不作为独立显著性检验。'
    content=SOURCE.read_text(encoding='utf-8')
    for k,v in {'MISS_ACC':f"{avg['accuracy']*100:.2f}",'MISS_F1':f"{avg['macro_f1']:.4f}",'MISS_MAE':f"{avg['mae']:.4f}",'MISS_CORR':f"{avg['pearson']:.4f}",'PREDICTION_TABLE':table,'MISSING_DISCUSSION':discussion}.items():content=content.replace('{{'+k+'}}',v)
    SOURCE.write_text(content,encoding='utf-8')
    shutil.copyfile(RUN/'attachment3_predictions.csv',SOURCE.parent/'问题二_C主线附件3历史预测.csv')
    shutil.copyfile(TEAM/'deployment_candidate_seed20260925/attachment3_predictions.csv',SOURCE.parent/'问题二_最终第三问入口模型_附件3预测.csv')
    (ASSETS/'validation_selected_snapshot.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
    sources=[RUN/'validation_grid.json',RUN/'attachment3_predictions.csv',TEAM/'MINILM_README.md',TEAM/'deployment_candidate_seed20260925/metadata.json',TEAM/'deployment_candidate_seed20260925/attachment3_predictions.csv']+list((ROOT/'docs/C/paper').glob('*论文*.md'))
    (ASSETS/'source_manifest.json').write_text(json.dumps({str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},ensure_ascii=False,indent=2),encoding='utf-8')
    print('ASSETS_OK',avg)

def expanded_assets():
    fig,ax=plt.subplots(figsize=(11,4));ax.axis('off');ax.set(xlim=(0,11),ylim=(0,4))
    boxes=[(.2,2.6,2.7,1,'历史基础框架\n掩码＋MiniLM–Fusion'),(4,2.6,3,1,'估计与目标：第1—3轮\n搜索／集成／任务权衡'),(4,.8,3,1,'表示与交互：第4—7轮\n容量／专家／微调／时序'),(8.3,1.7,2.5,1,'方法比较与边界分析\n另评估最终S0')]
    for x,y,w,h,t in boxes:ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.06',fc='#EDF2F6',ec='#355C7D'));ax.text(x+w/2,y+h/2,t,ha='center',va='center')
    for a,b in [((3,3.1),(3.9,3.1)),((1.6,2.5),(3.9,1.3)),((7.1,3.1),(8.2,2.4)),((7.1,1.3),(8.2,2.0))]:ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'->','color':'#355C7D'})
    ax.text(5.5,.15,'局部改进 → 独立核查 → 按实际推理路径评价；不把不同数据划分合并排名',ha='center');save(fig,'research_route')
    fig,aa=plt.subplots(1,2,figsize=(10,3.7))
    aa[0].bar(['原融合64维','加宽融合91维'],[60.30,61.31],color=['#355C7D','#6C9A8B']);aa[0].set(ylim=(58,64),ylabel='valid ACC（%）',title='第四轮：同轮三种子均值')
    for i,v in enumerate([60.30,61.31]):aa[0].text(i,v+.1,f'{v:.2f}',ha='center')
    x=np.arange(3);aa[1].bar(x-.17,[1.79,-.55,-.69],.34,label='第五轮：专家门控',color='#355C7D');aa[1].bar(x+.17,[-.274725,2.06044,2.1978],.34,label='第七轮：全valid校准',color='#D8933D');aa[1].axhline(0,color='black',lw=.7);aa[1].set(xticks=x,xticklabels=['配对种子1','配对种子2','配对种子3'],ylabel='相对同轮参照的ACC变化（百分点）');aa[1].legend(fontsize=8)
    for a in aa:a.grid(axis='y',alpha=.2)
    fig.tight_layout();save(fig,'controlled_gains')
    f=TEAM/'minilm_finetune/validation_grid_full_calibrated_seed20260925.json';grid=json.loads(f.read_text(encoding='utf-8'))['metrics'];clean=grid[0];missing=grid[1:];assert len(missing)==63
    predfile=TEAM/'minilm_finetune/validation_clean_predictions_calibrated_seed20260925.json';pred=json.loads(predfile.read_text(encoding='utf-8'));assert len(pred)==728
    cm=np.array(clean['confusion_matrix_true_by_pred']);fig,aa=plt.subplots(1,2,figsize=(10,4.1));aa[0].imshow(cm/cm.sum(1)[:,None],cmap='Blues',vmin=0,vmax=1)
    for i in range(3):
        for j in range(3):aa[0].text(j,i,f'{cm[i,j]}\n{cm[i,j]/cm[i].sum():.1%}',ha='center',va='center',color='white' if cm[i,j]/cm[i].sum()>.55 else 'black')
    aa[0].set(xticks=range(3),yticks=range(3),xticklabels=['负向','中性','正向'],yticklabels=['负向','中性','正向'],xlabel='预测类别',ylabel='真实类别')
    aa[1].scatter([r['target_score'] for r in pred],[r['served_score'] for r in pred],s=12,alpha=.24,color='#355C7D');aa[1].plot([-3,3],[-3,3],'--',color='#D8933D');aa[1].set(xlim=(-3.1,3.1),ylim=(-3.1,3.1),xlabel='真实强度',ylabel='一致性预测强度');fig.tight_layout();save(fig,'final_diagnostics')
    modes=list(dict.fromkeys(r['mode'] for r in missing));rates=[.1,.3,.5]
    def val(mode=None,rate=None,pos=None,key='macro_f1'):
        return float(np.mean([r[key] for r in missing if (mode is None or r['mode']==mode) and (rate is None or r['condition']['rate']==rate) and (pos is None or r['condition']['position']==pos)]))
    heat=np.array([[val(m,r) for r in rates] for m in modes]);fig,aa=plt.subplots(1,2,figsize=(10,4.4));aa[0].imshow(heat,cmap='Blues',vmin=.50,vmax=.62,aspect='auto');labels=[m.replace('text','文本').replace('audio','音频').replace('vision','视觉') for m in modes];aa[0].set(yticks=range(7),yticklabels=labels,xticks=range(3),xticklabels=['10%','30%','50%'],title='模态组合 × 长度：Macro-F1')
    for i in range(7):
        for j in range(3):aa[0].text(j,i,f'{heat[i,j]:.3f}',ha='center',va='center',color='white' if heat[i,j]>.575 else 'black')
    for pos,label,color in zip(['start','middle','end'],['起始','中部','末尾'],['#355C7D','#6C9A8B','#D8933D']):aa[1].plot([10,30,50],[100*val(rate=r,pos=pos,key='accuracy') for r in rates],'-o',label=label,color=color)
    aa[1].set(xlabel='缺失比例（%）',ylabel='ACC（%）',xticks=[10,30,50],title='位置 × 长度：准确率');aa[1].legend();aa[1].grid(alpha=.2);fig.tight_layout();save(fig,'final_missingness')
    desc=f"最终基线在63种条件下平均ACC为{100*val(key='accuracy'):.2f}%，Macro-F1为{val():.4f}，MAE为{val(key='mae'):.4f}，Pearson为{val(key='pearson'):.4f}；网格内最低Macro-F1为{min(r['macro_f1'] for r in missing):.4f}。10%、30%、50%缺失的平均ACC分别为"+'、'.join(f'{100*val(rate=r,key="accuracy"):.2f}%' for r in rates)+'。这些条件使用相同验证样本并等权汇总，不能把重复扰动当作新增独立样本；即使某些轻度缺失条件偶然优于原始输入，也不能解释为信息缺失本身必然有益。'
    s=SOURCE.read_text(encoding='utf-8').replace('{{FINAL_MISSING}}',desc)
    s=s.replace('采用八组缓存视图复用文本编码计算。','F₀采用八组缓存视图复用文本编码计算；B₀在线编码受损词元。')
    s=s.replace('附件3已归档结果对应种子20260924、第3轮的锁定检查点。','冻结配置选定种子20260924、第3轮检查点。')
    SOURCE.write_text(s,encoding='utf-8')
    manifest=json.loads((ASSETS/'source_manifest.json').read_text(encoding='utf-8'))
    for file in [f,predfile,SOURCE.parent/'问题一.docx']:manifest[str(file.relative_to(ROOT))]=hashlib.sha256(file.read_bytes()).hexdigest()
    (ASSETS/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')

def math_block(block):
    from latex2mathml.converter import convert
    from lxml import etree
    b=block.strip('$\n');match=re.search(r'(?:\\tag\{(5\.3-\d+)\}|\((5\.3-\d+)\))',b);number=next(x for x in match.groups() if x)
    body=b[:match.start()].strip().rstrip('。.')
    if '\\' not in body:
        if body.startswith('M̃'):body=r'\widetilde M_i^m(t)=M_i^m(t)[1-\mathbf1(m\in S)\mathbf1(t\in I_i)],\quad w_i=\min(L_i-1,\max(1,\mathrm{round}(\rho L_i)))'
        elif 'α' in body:body=r'\alpha_t^m=\frac{\widetilde M_t^m\exp(s_t^m)}{\sum_u\widetilde M_u^m\exp(s_u^m)},\qquad h^m=\sum_t\alpha_t^mH_t^m'
        elif 'MaskedSoftmax' in body:body=r'g=\mathrm{MaskedSoftmax}(W_g[h^T;h^A;h^V;q]+b_g),\quad h^f=\sum_mg^mh^m'
        elif 'SmoothL1' in body:body=r'L=L_{\mathrm{CE},w}+0.8L_{\mathrm{reg}}'
        elif 'mean条件' in body:body=r'S=\mathrm{mean}_{\omega}(\mathrm{MacroF1}_{\omega}-0.15\mathrm{MAE}_{\omega}+0.05r_{P,\omega})'
        else:raise ValueError(body)
    transform=etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
    result=transform(etree.fromstring(convert(body).encode('utf-8'))).getroot()
    math=result if result.tag==qn('m:oMath') else result.find('.//'+qn('m:oMath'))
    math.append(text('  ('+number+')'))
    return element('oMathPara',math)

def export():
    reference=Document(SOURCE.parent/'问题一.docx');ref_table=reference.tables[0]
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
    for name,size in [('Normal',10.5),('Title',16),('Heading 1',13),('Heading 2',11),('Caption',9)]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ('Normal','Caption') else '黑体')
    doc.styles['Normal'].paragraph_format.line_spacing=1.2;doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    for block in SOURCE.read_text(encoding='utf-8').strip().split('\n\n'):
        if block.startswith('# '):doc.add_heading(block[2:],0)
        elif block.startswith('## '):doc.add_heading(block[3:],1)
        elif block.startswith('### '):doc.add_heading(block[4:],2)
        elif block.startswith('$$'):
            p=doc.add_paragraph();p._p.append(math_block(block));p.paragraph_format.keep_together=True
        elif block.startswith('|'):
            rr=[[c.strip() for c in l.strip('|').split('|')] for l in block.splitlines()];rr=[rr[0]]+rr[2:];n=len(rr[0]);t=doc.add_table(rows=0,cols=n);t.style='Table Grid'
            props=deepcopy(ref_table._tbl.tblPr);props.find(qn('w:tblStyle')).set(qn('w:val'),'TableGrid');t._tbl.replace(t._tbl.tblPr,props)
            for i,row in enumerate(rr):
                cells=t.add_row().cells;pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if i==0:pr.append(OxmlElement('w:tblHeader'))
                for c,value in zip(cells,row):
                    ref=ref_table.cell(0 if i==0 else 1,0)
                    tc=c._tc.get_or_add_tcPr()
                    for name in ['vAlign','shd']:
                        el=ref._tc.tcPr.find(qn('w:'+name))
                        if el is not None:tc.append(deepcopy(el))
                    if i==0:tc.append(deepcopy(ref._tc.tcPr.find(qn('w:tcBorders'))))
                    p=c.paragraphs[0];p._p.insert(0,deepcopy(ref.paragraphs[0]._p.pPr));inline(p,value);p.paragraph_format.keep_with_next=i<len(rr)-1
                    for run in p.runs:
                        if run._r.rPr is not None:run._r.remove(run._r.rPr)
                        run._r.insert(0,deepcopy(ref.paragraphs[0].runs[0]._r.rPr))
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif block.startswith('!['):
            p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(SOURCE.parent/re.search(r'\]\((.+)\)',block).group(1)),width=Cm(10.5 if 'confusion.png' in block else 16.5))
        else:
            cap=bool(re.match(r'[图表]5-',block));p=doc.add_paragraph(style='Caption' if cap else 'Normal');inline(p,block)
            if block.startswith('表5-'):p.paragraph_format.keep_with_next=True
            if cap:p.paragraph_format.keep_together=True
            if block.startswith('表5-'):
                ref=next(p for p in reference.paragraphs if p.text.startswith('表5-1 '))
                old=p._p.pPr
                if old is not None:p._p.remove(old)
                p._p.insert(0,deepcopy(ref._p.pPr));p.paragraph_format.keep_with_next=True
                for run in p.runs:
                    if run._r.rPr is not None:run._r.remove(run._r.rPr)
                    run._r.insert(0,deepcopy(ref.runs[0]._r.rPr))
    p=sec.footer.paragraphs[0];p.alignment=1;f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    doc.core_properties.title='问题二的模型建立与求解';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(SOURCE.with_suffix('.docx'))
    has_final='## 5.3.11 ' in SOURCE.read_text(encoding='utf-8')
    expected=25 if has_final else 22 if '## 5.3.10 ' in SOURCE.read_text(encoding='utf-8') else 19
    assert len(doc.tables)==expected and len(doc.inline_shapes)==(9 if has_final else 8)
    assert '{{' not in SOURCE.read_text(encoding='utf-8')
    print('WORD_OK',len(doc.tables),len(doc.inline_shapes))

def audit():
    import win32com.client
    app=win32com.client.DispatchEx('Word.Application');app.Visible=False;app.DisplayAlerts=0;d=None
    try:
        d=app.Documents.Open(str(SOURCE.with_suffix('.docx')),ReadOnly=True,AddToRecentFiles=False);d.Repaginate()
        tables=[]
        for t in d.Tables:
            r=t.Range;pair=[d.Range(r.Start,r.Start).Information(3),d.Range(r.End-1,r.End-1).Information(3)]
            assert pair[0]==pair[1],pair
            tables.append(pair)
        figs=[x.Range.Information(3) for x in d.InlineShapes]
        caps=[p.Range.Information(3) for p in d.Paragraphs if re.match(r'图5-\d+ ',p.Range.Text)]
        assert figs==caps,(figs,caps)
        out=ROOT/'C/outputs/q2_integrated_paper_review';out.mkdir(exist_ok=True)
        d.ExportAsFixedFormat(str(out/'preview.pdf'),17)
        result=dict(pages=d.ComputeStatistics(2),tables=tables,figures=figs,captions=caps,equations=d.OMaths.Count,docx_sha256=hashlib.sha256(SOURCE.with_suffix('.docx').read_bytes()).hexdigest())
        (ASSETS/'word_layout_check.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print('LAYOUT_OK',result)
    finally:
        if d is not None:d.Close(0)
        app.Quit(0)

if __name__=='__main__':
    if '--check' in sys.argv:audit()
    else:
        assets();expanded_assets()
        from C.scripts.supplement_q2_paper import supplement
        supplement()
        completed=ROOT/'C/outputs/q2_three_layer_ablation_v2/summary.json'
        if completed.exists() and json.loads(completed.read_text(encoding='utf-8')).get('complete'):
            from C.scripts.report_q2_three_layer_ablation import report
            report()
        from C.scripts.integrate_q2_single_final import integrate
        integrate();export()
