"""Export the editorial Markdown to Word with embedded figures and native equations."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'docs/C/paper/第二轮集成优化实验论文草稿.md'


def element(tag,*children):
    node=OxmlElement('m:'+tag)
    for child in children: node.append(child)
    return node


def text(value):
    node=OxmlElement('m:t'); node.text=value
    return element('r',node)


def sub(value,index):
    return element('sSub',element('e',text(value)),element('sub',text(index)))


def power(value,index):
    return element('sSup',element('e',value),element('sup',text(index)))


def fraction(top,bottom):
    return element('f',element('num',*top),element('den',*bottom))


def summation(lower,upper,body):
    char=OxmlElement('m:chr'); char.set(qn('m:val'),'∑')
    loc=OxmlElement('m:limLoc'); loc.set(qn('m:val'),'undOvr')
    return element('nary',element('naryPr',char,loc),element('sub',text(lower)),
                   element('sup',text(upper)),element('e',*body))


def equation(block):
    if 'p_{si}' in block:
        terms=[sub('p','si'),text(' = '),summation('m=1','6',[sub('w','m'),sub('p','msi')]),
               text('，   '),sub('r','si'),text(' = '),summation('m=1','6',[sub('w','m'),sub('r','msi')]),
               text('，   '),sub('w','m'),text(' ≥ 0')]
    elif '\mathcal S' in block:
        terms=[text('S = '),fraction([text('1')],[text('4')]),summation('s∈𝒮','',[
            text('('),sub('MacroF1','s'),text(' − 0.15 '),sub('MAE','s'),text(' + 0.05 '),sub('Pearson','s'),text(')')])]
    elif 'J(w)' in block:
        terms=[text('J(w) = S(w) − 0.01 C(w)，   C(w) = '),fraction([
            text('6'),summation('m=1','6',[power(sub('w','m'),'2')]),text(' − 1')],[text('5')])]
    else: raise ValueError('Unsupported display equation')
    return element('oMathPara',element('oMath',*terms))


def inline(paragraph,value):
    # Editorial links retain their labels in the printable draft; the Markdown
    # remains the navigable source. Emphasis is preserved and code is readable.
    value=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',value)
    parts=re.split(r'(\*\*.*?\*\*|`[^`]+`|\$[^$]+\$)',value)
    for part in parts:
        if part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold=True
        elif part.startswith('`'):
            paragraph.add_run(part[1:-1])
        elif part.startswith('$'):
            formula=part[1:-1].replace('\\sum_m','∑ₘ').replace('\\geq','≥')
            formula=re.sub(r'_\{([^}]+)\}',r'_\1',formula)
            paragraph._p.append(element('oMath',text(formula)))
        elif part: paragraph.add_run(part)


def main():
    markdown=SOURCE.read_text(encoding='utf-8')
    doc=Document(); section=doc.sections[0]
    section.page_width,section.page_height=Cm(21),Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2.0)
    section.left_margin=section.right_margin=Cm(2.0)
    for name in ('Normal','Title','Heading 1','Heading 2','Caption'):
        style=doc.styles[name]
        style.font.name='Times New Roman'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ('Normal','Caption') else '黑体')
        style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.25
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',18),('Heading 1',14),('Heading 2',12),('Caption',9)]:
        doc.styles[name].font.size=Pt(size)
    doc.styles['Caption'].paragraph_format.line_spacing=1.1
    blocks=markdown.strip().split('\n\n')
    for block in blocks:
        if block.startswith('# '): doc.add_heading(block[2:],0)
        elif block.startswith('## '): doc.add_heading(block[3:],1)
        elif block.startswith('### '): doc.add_heading(block[4:],2)
        elif block.startswith('$$'):
            p=doc.add_paragraph(); p._p.append(equation(block)); p.paragraph_format.keep_together=True
        elif block.startswith('|'):
            rows=[[c.strip() for c in row.strip('|').split('|')] for row in block.splitlines()]
            rows=[rows[0]]+rows[2:]
            table=doc.add_table(rows=0,cols=len(rows[0])); table.style='Table Grid'
            for i,row in enumerate(rows):
                cells=table.add_row().cells
                props=table.rows[-1]._tr.get_or_add_trPr(); props.append(OxmlElement('w:cantSplit'))
                if i==0: props.append(OxmlElement('w:tblHeader'))
                for cell,value in zip(cells,row):
                    p=cell.paragraphs[0]; inline(p,value)
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.0
                    for run in p.runs: run.font.size=Pt(9); run.bold=i==0
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif block.startswith('!['):
            relative=re.search(r'\]\((.+)\)',block).group(1)
            p=doc.add_paragraph(); p.alignment=1
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(SOURCE.parent/relative),width=Cm(17))
        elif block.startswith('```'):
            p=doc.add_paragraph()
            p.paragraph_format.keep_together=True
            for line in block.splitlines()[1:-1]:
                run=p.add_run(line+'\n'); run.font.name='Consolas'; run.font.size=Pt(8)
        else:
            caption=block.startswith('**图C2-')
            p=doc.add_paragraph(style='Caption' if caption else 'Normal')
            inline(p,block)
            if caption: p.paragraph_format.keep_together=True
    footer=section.footer.paragraphs[0]; footer.alignment=1
    footer.add_run('第二轮实验临时稿 · ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.core_properties.title='第二轮集成优化实验论文草稿'
    doc.core_properties.subject='固定模型库集成搜索：实测结果、图表及局限'
    doc.core_properties.comments='Markdown为编辑源；含五幅实测配图，临时统稿片段。'
    target=SOURCE.with_suffix('.docx'); doc.save(target)
    check=Document(target)
    assert len(check.inline_shapes)==5 and len(check.tables)==2
    assert len(check._element.xpath('.//m:oMathPara'))==3
    assert not any('$$' in p.text or '```' in p.text or '![' in p.text for p in check.paragraphs)
    print('Exported Word: 5 embedded figures, 2 tables, 3 native display equations.')


if __name__=='__main__': main()
