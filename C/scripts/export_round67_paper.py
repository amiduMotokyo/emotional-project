"""Export editorial source to Word with native equations; optional Word pagination audit."""
import sys,re,json,hashlib,argparse
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT))
from C.scripts.export_round2_paper import inline,element,text,sub,summation
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
SOURCE=ROOT/'docs/C/paper/第六七轮优化局限与失败经验论文片段.md'
ASSETS=SOURCE.parent/'assets/round67'

def equation(block):
    if 'L=' in block:
        terms=[text('L = '),sub('L','CE,w'),text(' + 0.8 '),sub('L','reg'),text(' + 0.05 '),sub('L','SupCon')]
    else:
        terms=[text('h = '),summation('m∈{T,A,V}','',[sub('w','m'),sub('h','m')]),text('，  '),summation('m','',[sub('w','m')]),text(' = 1')]
    return element('oMathPara',element('oMath',*terms))

def export():
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
    for name,size in [('Normal',10.5),('Title',18),('Heading 1',14),('Heading 2',12),('Caption',9)]:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '黑体')
        style.paragraph_format.widow_control=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.25;doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    for block in SOURCE.read_text(encoding='utf-8').strip().split('\n\n'):
        if block.startswith('# '):doc.add_heading(block[2:],0)
        elif block.startswith('## '):doc.add_heading(block[3:],1)
        elif block.startswith('### '):doc.add_heading(block[4:],2)
        elif block.startswith('$$'):
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
            p=doc.add_paragraph();p._p.append(equation(block));p.paragraph_format.keep_together=True
        elif block.startswith('|'):
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
            rows=[[c.strip() for c in line.strip('|').split('|')] for line in block.splitlines()];rows=[rows[0]]+rows[2:]
            widths={5:[1.3,5.4,3.1,3.6,3.6],4:[7,3.3,3.3,3.4],3:[7,5,5]}[len(rows[0])]
            table=doc.add_table(rows=0,cols=len(widths));table.style='Table Grid';table.autofit=False
            for col,w in zip(table.columns,widths):col.width=Cm(w)
            for i,row in enumerate(rows):
                cells=table.add_row().cells;props=table.rows[-1]._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
                if i==0:props.append(OxmlElement('w:tblHeader'))
                for cell,value,w in zip(cells,row,widths):
                    cell.width=Cm(w);p=cell.paragraphs[0];inline(p,value);p.paragraph_format.line_spacing=1
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=i<len(rows)-1
                    for run in p.runs:run.font.size=Pt(9);run.bold=i==0
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif block.startswith('!['):
            p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
            path=SOURCE.parent/re.search(r'\]\((.+)\)',block).group(1);p.add_run().add_picture(str(path),width=Cm(17))
        else:
            caption=block.startswith('**图C67-')
            p=doc.add_paragraph(style='Caption' if caption else 'Normal');inline(p,block)
            if caption:p.paragraph_format.keep_together=True
    p=sec.footer.paragraphs[0];p.alignment=1;p.add_run('第六、七轮优化局限 · ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.core_properties.title='第六七轮优化局限与失败经验论文片段'
    doc.core_properties.comments='Markdown为唯一正文编辑源；图表来自已保存实测结果，未读取test。'
    doc.save(SOURCE.with_suffix('.docx'))
    check=Document(SOURCE.with_suffix('.docx'))
    assert len(check.tables)==4 and len(check.inline_shapes)==4 and len(check._element.xpath('.//m:oMathPara'))==2
    assert not any('$$' in p.text or '![' in p.text for p in check.paragraphs)
    print('WORD_EXPORT_OK: 4 figures, 4 tables, 2 editable equations',flush=True)

def audit():
    import win32com.client
    app=win32com.client.DispatchEx('Word.Application');app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3;doc=None
    try:
        doc=app.Documents.Open(str(SOURCE.with_suffix('.docx')),ReadOnly=True,AddToRecentFiles=False);doc.Repaginate()
        pics=[doc.InlineShapes(i).Range.Information(3) for i in range(1,doc.InlineShapes.Count+1)]
        caps=[p.Range.Information(3) for p in doc.Paragraphs if re.match('图C67-[0-9] ',p.Range.Text)]
        assert pics==caps and len(pics)==4,(pics,caps)
        tables=[]
        for i in range(1,doc.Tables.Count+1):
            r=doc.Tables(i).Range;start=doc.Range(r.Start,r.Start).Information(3);end=doc.Range(r.End-1,r.End-1).Information(3)
            assert start==end,(i,start,end);tables.append(start)
        assert doc.OMaths.Count==2
        preview=ROOT/'C/outputs/round67_paper_preview';preview.mkdir(parents=True,exist_ok=True)
        doc.ExportAsFixedFormat(str(preview/'preview.pdf'),17)
        result=dict(pages=doc.ComputeStatistics(2),figure_pages=pics,caption_pages=caps,table_pages=tables,equations=doc.OMaths.Count,
                    docx_sha256=hashlib.sha256(SOURCE.with_suffix('.docx').read_bytes()).hexdigest())
        (ASSETS/'word_layout_check.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print('WORD_LAYOUT_OK',result,flush=True)
    finally:
        if doc is not None:doc.Close(0)
        app.Quit(0)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--check',action='store_true');args=parser.parse_args()
    export()
    if args.check:audit()
