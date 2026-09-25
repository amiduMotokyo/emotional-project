"""Export third-round Markdown with embedded figures and editable Word equations."""
import sys
import re
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from C.scripts.export_round2_paper import element, text, sub, fraction, summation, inline
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'docs/C/paper/第三轮准确率优化实验论文草稿.md'

def equation(block):
    if 'L=' in block:
        terms=[text('L = −'),fraction([summation('i=1','N',[sub('w','yᵢ'),text(' log '),sub('p','i,yᵢ')])],
               [summation('i=1','N',[sub('w','yᵢ')])]),text(' + '),fraction([text('λ')],[text('N')]),
               summation('i=1','N',[text('H('),sub('r','i'),text(' − '),sub('s','i'),text(')')])]
    elif '\\tilde' in block:
        terms=[sub('p̃','ic'),text(' = '),fraction([sub('p̄','ic'),text(' exp('),sub('b','c'),text(')')],
               [summation('k=0','2',[sub('p̄','ik'),text(' exp('),sub('b','k'),text(')')])]),text('，  '),sub('b','2'),text(' = 0')]
    else:
        terms=[sub('p̄','i'),text(' = '),fraction([text('1')],[text('M')]),summation('m=1','M',[sub('p','mi')]),
               text('，  '),sub('r̄','i'),text(' = '),fraction([text('1')],[text('M')]),summation('m=1','M',[sub('r','mi')])]
    return element('oMathPara',element('oMath',*terms))

def main():
    doc=Document();sec=doc.sections[0]
    sec.page_width,sec.page_height=Cm(21),Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
    for name,size in [('Normal',10.5),('Title',18),('Heading 1',14),('Heading 2',12),('Caption',9)]:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ('Normal','Caption') else '黑体')
        style.paragraph_format.widow_control=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.25
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Caption'].paragraph_format.line_spacing=1.1
    for block in SOURCE.read_text(encoding='utf-8').strip().split('\n\n'):
        if block.startswith('# '): doc.add_heading(block[2:],0)
        elif block.startswith('## '):doc.add_heading(block[3:],1)
        elif block.startswith('### '):doc.add_heading(block[4:],2)
        elif block.startswith('$$'):
            p=doc.add_paragraph();p._p.append(equation(block));p.paragraph_format.keep_together=True
        elif block.startswith('|'):
            rows=[[c.strip() for c in line.strip('|').split('|')] for line in block.splitlines()]
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid';table.autofit=False
            widths=[2.8,2.8,2.5,2.5,6.4] if len(doc.tables)==1 else [5,3,3,3,3]
            for col,width in zip(table.columns,widths):col.width=Cm(width)
            for i,row in enumerate([rows[0]]+rows[2:]):
                cells=table.add_row().cells;props=table.rows[-1]._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
                if i==0:props.append(OxmlElement('w:tblHeader'))
                for cell,value,width in zip(cells,row,widths):
                    cell.width=Cm(width)
                    p=cell.paragraphs[0];inline(p,value);p.paragraph_format.line_spacing=1
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.keep_with_next=i<len(rows)-2
                    for r in p.runs:r.font.size=Pt(9);r.bold=i==0
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif block.startswith('!['):
            p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
            path=SOURCE.parent/re.search(r'\]\((.+)\)',block).group(1)
            # Compact square confusion matrix; wide multi-panel figures use the text width.
            p.add_run().add_picture(str(path),width=Cm(12 if 'fig5_' in path.name else 17))
        else:
            caption=block.startswith('**图C3-')
            p=doc.add_paragraph(style='Caption' if caption else 'Normal');inline(p,block)
            if caption:p.paragraph_format.keep_together=True
    p=sec.footer.paragraphs[0];p.alignment=1;p.add_run('第三轮实验临时稿 · ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    doc.core_properties.title='第三轮准确率优化实验论文草稿'
    doc.core_properties.subject='训练消融、局部微调、校准与负结果分析'
    doc.core_properties.comments='Markdown为编辑源；实测图表与原生可编辑公式。'
    target=SOURCE.with_suffix('.docx');doc.save(target)
    check=Document(target)
    assert len(check.inline_shapes)==5 and len(check.tables)==2
    assert len(check._element.xpath('.//m:oMathPara'))==3
    assert not any('$$' in p.text or '![' in p.text for p in check.paragraphs)
    print('Exported third-round Word: 5 figures, 2 tables, 3 native equations.')

if __name__=='__main__':main()
