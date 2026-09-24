from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml


def load_config(config_path: str | Path | None = None) -> dict[str, Any]:
    if config_path is None:
        project_root = Path(__file__).resolve().parents[1]
        config_path = project_root / "configs" / "config.yaml"
    else:
        config_path = Path(config_path).resolve()

    with config_path.open("r", encoding="utf-8") as handle:
        cfg = yaml.safe_load(handle)

    cfg["_config_path"] = str(config_path)
    cfg["_project_root"] = str(config_path.parent.parent)
    return cfg


def resolve_path(value: str | Path, project_root: str | Path) -> Path:
    path = Path(value)
    if path.is_absolute():
        return path
    return Path(project_root) / path


def resolve_config_paths(cfg: dict[str, Any]) -> dict[str, Path]:
    root = Path(cfg["_project_root"])
    paths = cfg["paths"]
    ffmpeg_value = Path(paths["ffmpeg_exe"])
    if ffmpeg_value.is_absolute() or ffmpeg_value.parent != Path("."):
        ffmpeg_path = resolve_path(ffmpeg_value, root)
    else:
        local_candidate = root / ffmpeg_value
        ffmpeg_path = local_candidate if local_candidate.exists() else ffmpeg_value
    resolved = {
        "project_root": root,
        "video_root": resolve_path(paths["video_root"], root),
        "audio_root": resolve_path(paths["audio_root"], root),
        "label_file": resolve_path(paths["label_file"], root),
        "bert_path": resolve_path(paths["bert_path"], root),
        "openface_exe": resolve_path(paths["openface_exe"], root),
        "ffmpeg_exe": ffmpeg_path,
        "ffprobe_exe": resolve_path(paths["ffprobe_exe"], root),
        "output_root": resolve_path(paths["output_root"], root),
        "work_root": resolve_path(paths["work_root"], root),
    }
    return resolved


def prepare_output_dirs(cfg: dict[str, Any]) -> dict[str, Path]:
    paths = resolve_config_paths(cfg)
    output_root = paths["output_root"]
    dirs = {
        "features": output_root / "features",
        "audio": output_root / "features" / "audio",
        "visual": output_root / "features" / "visual",
        "text": output_root / "features" / "text",
        "alignment": output_root / "alignment",
        "summary": output_root / "summary",
        "logs": output_root / "logs",
        "evidence": output_root / "evidence",
        "work": paths["work_root"],
    }
    for directory in dirs.values():
        directory.mkdir(parents=True, exist_ok=True)
    return dirs


def normalize_identifier(value: Any) -> str:
    if pd.isna(value):
        raise ValueError("Sample identifier cannot be missing.")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def load_labels(cfg: dict[str, Any]) -> pd.DataFrame:
    paths = resolve_config_paths(cfg)
    if not paths["label_file"].exists():
        raise FileNotFoundError(f"Label file not found: {paths['label_file']}")

    try:
        labels = pd.read_excel(paths["label_file"], sheet_name="label")
    except (ValueError, ImportError):
        import openpyxl

        workbook = openpyxl.load_workbook(
            paths["label_file"], read_only=True, data_only=True
        )
        sheet = workbook["label"] if "label" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(value).strip() for value in next(rows)]
        labels = pd.DataFrame(rows, columns=header)
        workbook.close()

    required = {"video_id", "clip_id", "text", "label", "annotation"}
    missing = required.difference(labels.columns)
    if missing:
        raise ValueError(f"Missing label columns: {sorted(missing)}")

    labels = labels.copy()
    labels["video_id"] = labels["video_id"].map(normalize_identifier)
    labels["clip_id"] = labels["clip_id"].map(normalize_identifier)
    labels["sample_id"] = labels["video_id"] + "_" + labels["clip_id"]
    labels["text"] = labels["text"].fillna("").astype(str)
    return labels


def video_path(video_root: Path, video_id: str, clip_id: str) -> Path:
    return video_root / video_id / f"{clip_id}.mp4"


def align_frame_index(
    features: np.ndarray,
    frame_shift_seconds: float,
    segment_len: float,
    max_segments: int,
) -> tuple[np.ndarray, int]:
    if features.ndim != 2:
        raise ValueError(f"Expected a 2-D feature array, got {features.shape}")

    aligned = np.zeros((max_segments, features.shape[1]), dtype=np.float32)
    counts = np.zeros(max_segments, dtype=np.int32)
    frame_segments = np.floor(
        (np.arange(features.shape[0], dtype=np.float64) * frame_shift_seconds)
        / segment_len
    ).astype(np.int64)
    valid_frames = frame_segments < max_segments

    np.add.at(aligned, frame_segments[valid_frames], features[valid_frames])
    np.add.at(counts, frame_segments[valid_frames], 1)
    nonzero = counts > 0
    aligned[nonzero] /= counts[nonzero, None]
    return aligned, int(nonzero.sum())


def align_timestamps(
    timestamps: np.ndarray,
    features: np.ndarray,
    segment_len: float,
    max_segments: int,
) -> tuple[np.ndarray, int]:
    if features.ndim != 2:
        raise ValueError(f"Expected a 2-D feature array, got {features.shape}")
    if len(timestamps) != len(features):
        raise ValueError("Timestamp count and feature row count do not match.")

    aligned = np.zeros((max_segments, features.shape[1]), dtype=np.float32)
    valid = np.zeros(max_segments, dtype=bool)
    for index in range(max_segments):
        mask = (timestamps >= index * segment_len) & (
            timestamps < (index + 1) * segment_len
        )
        if mask.any():
            aligned[index] = features[mask].mean(axis=0)
            valid[index] = True
    return aligned, int(valid.sum())


def write_json(path: str | Path, value: Any) -> None:
    with Path(path).open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
