from __future__ import annotations

import argparse
import json
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd


class LocalBertWordPieceTokenizer:
    def __init__(self, model_dir: str | Path):
        model_path = Path(model_dir)
        vocab_path = model_path / "vocab.txt"
        tokenizer_json = model_path / "tokenizer.json"
        if vocab_path.exists():
            tokens = vocab_path.read_text(encoding="utf-8").splitlines()
            self.vocab = {token: index for index, token in enumerate(tokens)}
        elif tokenizer_json.exists():
            payload = json.loads(tokenizer_json.read_text(encoding="utf-8"))
            self.vocab = {
                str(token): int(index)
                for token, index in payload["model"]["vocab"].items()
            }
        else:
            raise FileNotFoundError(
                f"Neither vocab.txt nor tokenizer.json exists in {model_path}"
            )
        self.unk_token = "[UNK]"
        self.max_input_chars_per_word = 100
        self._cache: dict[str, list[str]] = {}

    @staticmethod
    def _is_cjk(character: str) -> bool:
        codepoint = ord(character)
        return (
            0x4E00 <= codepoint <= 0x9FFF
            or 0x3400 <= codepoint <= 0x4DBF
            or 0x20000 <= codepoint <= 0x2A6DF
            or 0x2A700 <= codepoint <= 0x2B73F
            or 0x2B740 <= codepoint <= 0x2B81F
        )

    def _basic_tokens(self, text: str) -> list[str]:
        normalized = unicodedata.normalize("NFD", text)
        normalized = "".join(
            character
            for character in normalized
            if unicodedata.category(character) != "Mn"
        )
        normalized = normalized.lower()
        output: list[str] = []
        for character in normalized:
            if character == "\x00" or character == "\ufffd":
                continue
            category = unicodedata.category(character)
            if category.startswith("C") or category in {"Zs", "Zl", "Zp"}:
                output.append(" ")
            elif category.startswith("P") or self._is_cjk(character):
                output.extend([" ", character, " "])
            else:
                output.append(character)
        return "".join(output).strip().split()

    def _wordpiece(self, token: str) -> list[str]:
        if token in self._cache:
            return self._cache[token]
        if len(token) > self.max_input_chars_per_word:
            self._cache[token] = [self.unk_token]
            return self._cache[token]

        characters = list(token)
        output: list[str] = []
        start = 0
        while start < len(characters):
            end = len(characters)
            current = None
            while start < end:
                fragment = "".join(characters[start:end])
                candidate = fragment if start == 0 else f"##{fragment}"
                if candidate in self.vocab:
                    current = candidate
                    break
                end -= 1
            if current is None:
                self._cache[token] = [self.unk_token]
                return self._cache[token]
            output.append(current)
            start = end
        self._cache[token] = output
        return output

    def tokenize(self, text: str) -> list[str]:
        output: list[str] = []
        for token in self._basic_tokens(text):
            output.extend(self._wordpiece(token))
        return output


def build_tokenizer(model_dir: str | Path):
    try:
        from transformers import BertTokenizer

        return BertTokenizer.from_pretrained(
            str(model_dir),
            local_files_only=True,
        )
    except ImportError:
        return LocalBertWordPieceTokenizer(model_dir)


def normalize_word(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def align_reference_words(
    reference_words: list[str],
    whisper_words: list[dict],
) -> tuple[list[dict | None], int, int]:
    reference_normalized = [normalize_word(word) for word in reference_words]
    whisper_normalized = [
        normalize_word(str(word.get("word", ""))) for word in whisper_words
    ]
    matcher = SequenceMatcher(
        None,
        reference_normalized,
        whisper_normalized,
        autojunk=False,
    )
    assigned: list[dict | None] = [None] * len(reference_words)
    matched = 0
    extra_whisper = 0
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for offset in range(i2 - i1):
                assigned[i1 + offset] = whisper_words[j1 + offset]
                matched += 1
        elif tag == "replace":
            paired = min(i2 - i1, j2 - j1)
            for offset in range(paired):
                assigned[i1 + offset] = whisper_words[j1 + offset]
                matched += 1
            extra_whisper += max(0, (j2 - j1) - paired)
        elif tag == "insert":
            extra_whisper += j2 - j1
    return assigned, matched, extra_whisper


def build_token_pairs(
    reference_words: list[str],
    raw_text: str,
    tokenizer,
    max_length: int,
) -> list[tuple[str, int]]:
    flat_tokens: list[str] = []
    flat_word_ids: list[int] = []
    for word_index, word in enumerate(reference_words):
        for subword in tokenizer.tokenize(word):
            flat_tokens.append(subword)
            flat_word_ids.append(word_index)

    full_tokens = tokenizer.tokenize(str(raw_text))
    assigned_word_ids = [-1] * len(full_tokens)
    matcher = SequenceMatcher(None, flat_tokens, full_tokens, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for offset in range(i2 - i1):
                assigned_word_ids[j1 + offset] = flat_word_ids[i1 + offset]
        elif tag == "replace":
            paired = min(i2 - i1, j2 - j1)
            for offset in range(paired):
                assigned_word_ids[j1 + offset] = flat_word_ids[i1 + offset]
            for offset in range(paired, j2 - j1):
                source_index = min(i1 + offset, len(flat_word_ids) - 1)
                assigned_word_ids[j1 + offset] = (
                    flat_word_ids[source_index] if flat_word_ids else -1
                )
        elif tag == "insert":
            for offset in range(j2 - j1):
                source_index = min(j1 + offset, len(flat_word_ids) - 1)
                assigned_word_ids[j1 + offset] = (
                    flat_word_ids[source_index] if flat_word_ids else -1
                )

    pairs = [("[CLS]", -1)]
    pairs.extend(zip(full_tokens, assigned_word_ids))
    pairs.append(("[SEP]", -1))
    return pairs[:max_length]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--words-dir", required=True)
    parser.add_argument("--label-file", required=True)
    parser.add_argument("--bert-path", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--text-summary", default=None)
    parser.add_argument("--max-length", type=int, default=50)
    args = parser.parse_args()

    words_dir = Path(args.words_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    tokenizer = build_tokenizer(args.bert_path)

    try:
        labels = pd.read_excel(args.label_file, sheet_name="label")
    except Exception:
        import openpyxl

        workbook = openpyxl.load_workbook(
            args.label_file, read_only=True, data_only=True
        )
        sheet = workbook["label"] if "label" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(value).strip() for value in next(rows)]
        labels = pd.DataFrame(rows, columns=header)
        workbook.close()

    labels["sample_id"] = (
        labels["video_id"].astype(str).str.replace(r"\.0$", "", regex=True)
        + "_"
        + labels["clip_id"].astype(str).str.replace(r"\.0$", "", regex=True)
    )
    valid_token_map = {}
    if args.text_summary:
        text_summary = pd.read_csv(args.text_summary)
        valid_token_map = dict(
            zip(text_summary["sample_id"], text_summary["valid_tokens"])
        )
    rows = []
    coverage_rows = []
    for label_row in labels.itertuples(index=False):
        words_path = words_dir / f"{label_row.sample_id}_words.json"
        whisper_words = (
            json.loads(words_path.read_text(encoding="utf-8"))
            if words_path.exists()
            else []
        )
        reference_words = str(label_row.text).split()
        assigned, matched, extra_whisper = align_reference_words(
            reference_words,
            whisper_words,
        )

        token_pairs = build_token_pairs(
            reference_words,
            str(label_row.text),
            tokenizer,
            args.max_length,
        )
        expected_valid_tokens = int(
            valid_token_map.get(label_row.sample_id, len(token_pairs))
        )
        if expected_valid_tokens > 0 and len(token_pairs) > expected_valid_tokens:
            token_pairs = token_pairs[:expected_valid_tokens]
        real_token_count = sum(1 for _, word_index in token_pairs if word_index >= 0)
        tokens_with_timestamp = 0
        for token_position in range(args.max_length):
            if token_position < len(token_pairs):
                subword, word_index = token_pairs[token_position]
                token_valid = True
                source_word = (
                    reference_words[word_index] if word_index >= 0 else ""
                )
                word_timestamp = (
                    assigned[word_index] if word_index >= 0 else None
                )
                timestamp_available = word_timestamp is not None
                if word_index >= 0 and timestamp_available:
                    tokens_with_timestamp += 1
                rows.append(
                    {
                        "sample_id": label_row.sample_id,
                        "token_position": token_position,
                        "subword_token": subword,
                        "token_valid": token_valid,
                        "source_word_index": word_index,
                        "source_word": source_word,
                        "start_sec": (
                            float(word_timestamp.get("start", 0.0))
                            if timestamp_available
                            else None
                        ),
                        "end_sec": (
                            float(word_timestamp.get("end", 0.0))
                            if timestamp_available
                            else None
                        ),
                        "probability": (
                            float(word_timestamp.get("probability", 0.0))
                            if timestamp_available
                            else None
                        ),
                        "timestamp_available": timestamp_available,
                    }
                )
            else:
                rows.append(
                    {
                        "sample_id": label_row.sample_id,
                        "token_position": token_position,
                        "subword_token": "[PAD]",
                        "token_valid": False,
                        "source_word_index": -1,
                        "source_word": "",
                        "start_sec": None,
                        "end_sec": None,
                        "probability": None,
                        "timestamp_available": False,
                    }
                )
        coverage_rows.append(
            {
                "sample_id": label_row.sample_id,
                "reference_words": len(reference_words),
                "whisper_words": len(whisper_words),
                "matched_reference_words": matched,
                "unmatched_reference_words": len(reference_words) - matched,
                "extra_whisper_words": extra_whisper,
                "word_match_rate": matched / max(len(reference_words), 1),
                "word_coverage_capped": min(
                    matched / max(len(reference_words), 1),
                    1.0,
                ),
                "bert_valid_tokens_expected": expected_valid_tokens,
                "bert_tokens_in_mapping": len(token_pairs),
                "tokens_with_timestamp": tokens_with_timestamp,
                "token_timestamp_coverage": (
                    tokens_with_timestamp / max(real_token_count, 1)
                ),
            }
        )
    pd.DataFrame(rows).to_csv(
        output_dir / "text_time_alignment.csv",
        index=False,
        encoding="utf-8-sig",
    )
    coverage = pd.DataFrame(coverage_rows)
    coverage.to_csv(
        output_dir / "text_time_coverage.csv",
        index=False,
        encoding="utf-8-sig",
    )
    coverage[coverage["token_timestamp_coverage"] < 0.8].sort_values(
        ["token_timestamp_coverage", "sample_id"]
    ).to_csv(
        output_dir / "text_time_review.csv",
        index=False,
        encoding="utf-8-sig",
    )
    print(coverage.describe().to_string())


if __name__ == "__main__":
    main()
